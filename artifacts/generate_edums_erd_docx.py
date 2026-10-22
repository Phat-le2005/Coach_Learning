from __future__ import annotations

import html
import struct
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "artifacts" / "EduMS_ERD_Tong_Quat.docx"

GREEN = "1E5B43"
MINT = "DFF3E8"
CREAM = "FFF9ED"
VIOLET = "6A5AA3"
AMBER = "B7791F"
SKY = "DDEFF7"


def node(name: str, title: str, fields: list[str], color: str = MINT) -> str:
    rows = "".join(
        f'<TR><TD ALIGN="LEFT"><FONT POINT-SIZE="9">{html.escape(field)}</FONT></TD></TR>'
        for field in fields
    )
    return (
        f'"{name}" [shape=plain label=<<TABLE BORDER="1" CELLBORDER="0" '
        f'CELLSPACING="0" CELLPADDING="6" COLOR="#{GREEN}">'
        f'<TR><TD BGCOLOR="#{color}"><B>{html.escape(title)}</B></TD></TR>'
        f'{rows}</TABLE>>];'
    )


def graph(title: str, nodes: list[str], edges: list[tuple[str, str, str]]) -> str:
    edge_lines = [
        f'"{source}" -> "{target}" [label="{html.escape(label)}"];'
        for source, target, label in edges
    ]
    return "\n".join(
        [
            "digraph ERD {",
            'graph [rankdir=LR, bgcolor="transparent", pad="0.25", nodesep="0.35", ranksep="0.65", splines=ortho, labeljust=l];',
            f'label="{html.escape(title)}";',
            'labelloc="t";',
            'fontname="Arial";',
            'fontsize="18";',
            f'fontcolor="#{GREEN}";',
            'node [fontname="Arial", fontsize="10"];',
            f'edge [fontname="Arial", fontsize="8", color="#{GREEN}", fontcolor="#{GREEN}", arrowsize="0.7"];',
            *nodes,
            *edge_lines,
            "}",
        ]
    )


DIAGRAMS: list[tuple[str, str]] = [
    (
        "01_identity_classroom",
        graph(
            "Identity, RBAC, Profile và Classroom",
            [
                node("users", "users", ["PK id", "UK email", "password_hash", "full_name", "role", "is_active"]),
                node("auth_sessions", "auth_sessions", ["PK id", "FK user_id", "expires_at", "revoked_at"], SKY),
                node("refresh", "auth_refresh_tokens", ["PK id", "FK auth_session_id", "token_hash", "expires_at"], SKY),
                node("student", "student_profiles", ["PK id", "FK/UK user_id", "grade_level", "school_name", "is_synthetic"]),
                node("teacher", "teacher_profiles", ["PK id", "FK/UK user_id", "specialization"]),
                node("parent", "parent_profiles", ["PK id", "FK/UK user_id"]),
                node("parent_student", "parent_student_relationships", ["PK/FK parent_id", "PK/FK student_id", "relationship_type", "is_active"], CREAM),
                node("classroom", "classrooms", ["PK id", "name", "academic_year", "grade_level", "FK created_by"], CREAM),
                node("teacher_class", "teacher_classrooms", ["PK/FK teacher_id", "PK/FK classroom_id", "role"], CREAM),
                node("enrollment", "student_enrollments", ["PK/FK student_id", "PK/FK classroom_id", "status", "joined_at"], CREAM),
            ],
            [
                ("users", "auth_sessions", "1 : N"),
                ("auth_sessions", "refresh", "1 : N"),
                ("users", "student", "1 : 0..1"),
                ("users", "teacher", "1 : 0..1"),
                ("users", "parent", "1 : 0..1"),
                ("parent", "parent_student", "1 : N"),
                ("student", "parent_student", "1 : N"),
                ("users", "classroom", "creates"),
                ("teacher", "teacher_class", "1 : N"),
                ("classroom", "teacher_class", "1 : N"),
                ("student", "enrollment", "1 : N"),
                ("classroom", "enrollment", "1 : N"),
            ],
        ).replace("rankdir=LR", "rankdir=TB"),
    ),
    (
        "02_curriculum_questions",
        graph(
            "Curriculum, Knowledge Graph và Question Bank",
            [
                node("curriculum", "curricula", ["PK id", "UK code", "name"], CREAM),
                node("version", "curriculum_versions", ["PK id", "FK curriculum_id", "version", "status"], CREAM),
                node("subject", "subjects", ["PK id", "FK curriculum_version_id", "UK code", "name", "grade_level"]),
                node("chapter", "chapters", ["PK id", "FK subject_id", "code", "name", "order_index"]),
                node("topic", "topics", ["PK id", "FK chapter_id", "code", "name", "order_index"]),
                node("pt", "problem_types", ["PK id", "FK topic_id", "code", "name", "initial_difficulty", "is_active"]),
                node("concept", "concepts", ["PK id", "FK subject_id", "code", "name", "grade_level"], SKY),
                node("relation", "concept_relations", ["PK id", "FK source_concept_id", "FK target_concept_id", "relation_type"], SKY),
                node("lo", "learning_outcomes", ["PK id", "FK subject_id", "code", "description"], SKY),
                node("ptc", "problem_type_concepts", ["PK/FK problem_type_id", "PK/FK concept_id", "weight", "is_primary"], CREAM),
                node("ptlo", "problem_type_learning_outcomes", ["PK/FK problem_type_id", "PK/FK learning_outcome_id", "weight"], CREAM),
                node("question", "questions", ["PK id", "FK subject_id", "content_hash", "question_type", "difficulty", "correct_answer", "exam_weight", "status"], "F0EAF9"),
                node("option", "question_options", ["PK id", "FK question_id", "label", "content", "order_index"], "F0EAF9"),
                node("qc", "question_concepts", ["PK/FK question_id", "PK/FK concept_id", "weight"], "F0EAF9"),
                node("qpt", "question_problem_types", ["PK/FK question_id", "PK/FK problem_type_id", "weight", "is_primary"], "F0EAF9"),
                node("qlo", "question_learning_outcomes", ["PK/FK question_id", "PK/FK learning_outcome_id", "weight"], "F0EAF9"),
                node("group", "question_groups", ["PK id", "FK subject_id", "content", "group_type", "scoring_policy", "status"], "FDECC8"),
                node("group_item", "question_group_items", ["PK id", "FK group_id", "FK/UK question_id", "label a..d", "order_index 0..3"], "FDECC8"),
            ],
            [
                ("curriculum", "version", "1 : N"),
                ("version", "subject", "1 : N"),
                ("subject", "chapter", "1 : N"),
                ("chapter", "topic", "1 : N"),
                ("topic", "pt", "1 : N"),
                ("subject", "concept", "1 : N"),
                ("concept", "relation", "source/target"),
                ("subject", "lo", "1 : N"),
                ("pt", "ptc", "1 : N"),
                ("concept", "ptc", "1 : N"),
                ("pt", "ptlo", "1 : N"),
                ("lo", "ptlo", "1 : N"),
                ("subject", "question", "1 : N"),
                ("question", "option", "1 : N"),
                ("question", "qc", "1 : N"),
                ("concept", "qc", "1 : N"),
                ("question", "qpt", "1 : N"),
                ("pt", "qpt", "1 : N"),
                ("question", "qlo", "1 : N"),
                ("lo", "qlo", "1 : N"),
                ("subject", "group", "1 : N"),
                ("group", "group_item", "1 : 4"),
                ("question", "group_item", "1 : 0..1"),
            ],
        ),
    ),
    (
        "03_learning_mastery",
        graph(
            "Practice, Attempt, Evidence, Mastery và Digital Twin",
            [
                node("student", "student_profiles", ["PK id", "FK user_id"]),
                node("subject", "subjects", ["PK id", "code"]),
                node("session", "practice_sessions", ["PK id", "FK student_id", "FK subject_id", "session_type", "status", "started_at", "completed_at"], CREAM),
                node("question", "questions", ["PK id", "question_type", "difficulty", "status"], "F0EAF9"),
                node("attempt", "question_attempts", ["PK id", "FK session_id + student_id", "FK question_id", "student_answer", "is_correct", "score", "hint_count", "evaluation_status"], "F0EAF9"),
                node("evidence", "learning_evidence_events", ["PK id", "FK student_id", "FK subject_id", "source_kind", "source_entity_id", "is_correct", "score", "difficulty", "hint_count"], SKY),
                node("evc", "learning_evidence_concepts", ["PK/FK event_id", "PK/FK concept_id", "weight"], SKY),
                node("evp", "learning_evidence_problem_types", ["PK/FK event_id", "PK/FK problem_type_id", "weight"], SKY),
                node("concept", "concepts", ["PK id", "code", "name"]),
                node("pt", "problem_types", ["PK id", "code", "name"]),
                node("cm", "student_concept_mastery", ["PK/FK student_id", "PK/FK concept_id", "mastery_score", "confidence", "attempt_count", "correct_count"], "DDF4D9"),
                node("pm", "student_problem_type_mastery", ["PK/FK student_id", "PK/FK problem_type_id", "mastery_score", "confidence", "attempt_count", "correct_count"], "DDF4D9"),
                node("review", "review_schedules", ["PK id", "FK student_id", "FK concept_id?", "FK problem_type_id?", "next_review_at"], "FDECC8"),
                node("twin", "StudentDigitalTwin", ["READ MODEL - no table", "profile + mastery + attempts", "review + goals + risk signals"], "F6EAF2"),
            ],
            [
                ("student", "session", "1 : N"),
                ("subject", "session", "1 : N"),
                ("session", "attempt", "1 : N"),
                ("question", "attempt", "1 : N"),
                ("student", "evidence", "1 : N"),
                ("subject", "evidence", "1 : N"),
                ("evidence", "evc", "1 : N"),
                ("concept", "evc", "1 : N"),
                ("evidence", "evp", "1 : N"),
                ("pt", "evp", "1 : N"),
                ("student", "cm", "1 : N"),
                ("concept", "cm", "1 : N"),
                ("student", "pm", "1 : N"),
                ("pt", "pm", "1 : N"),
                ("student", "review", "1 : N"),
                ("concept", "review", "optional target"),
                ("pt", "review", "optional target"),
                ("cm", "twin", "aggregates"),
                ("pm", "twin", "aggregates"),
                ("attempt", "twin", "recent activity"),
            ],
        ).replace("rankdir=LR", "rankdir=TB"),
    ),
    (
        "04_diagnosis_repair",
        graph(
            "Diagnostic, Misconception, Root Cause và Adaptive Repair",
            [
                node("session", "practice_sessions", ["PK id", "FK student_id", "session_type"]),
                node("diagnostic", "diagnostic_assessments", ["PK/FK practice_session_id", "current_problem_type_id?", "current_concept_id?", "status"], CREAM),
                node("blueprint", "diagnostic_blueprint_items", ["PK id", "FK assessment_id", "FK problem_type_id", "importance"], CREAM),
                node("definition", "misconception_definitions", ["PK id", "UK code", "FK concept_id", "FK problem_type_id?", "category"], SKY),
                node("student_mis", "student_misconceptions", ["PK/FK student_id", "PK/FK definition_id", "confidence", "occurrence_count", "status"], SKY),
                node("evidence", "learning_evidence_events", ["PK id", "FK student_id", "source_kind", "score"]),
                node("diagnosis", "root_cause_diagnoses", ["PK id", "FK evidence_event_id", "FK student_id", "FK concept/problem_type", "FK misconception_id", "status"], "FDECC8"),
                node("candidate", "root_cause_candidates", ["PK id", "FK diagnosis_id", "candidate_kind", "FK concept/problem_type", "confidence"], "FDECC8"),
                node("verification", "root_cause_verifications", ["PK id", "FK diagnosis_id", "FK practice_session_id", "status"], "FDECC8"),
                node("repair", "adaptive_repair_sessions", ["PK id", "FK student_id", "FK diagnosis_id", "FK target concept/PT", "status"], "F0EAF9"),
                node("repair_item", "adaptive_repair_items", ["PK id", "FK repair_session_id", "FK question_id", "FK question_attempt_id?", "order_index"], "F0EAF9"),
            ],
            [
                ("session", "diagnostic", "1 : 0..1"),
                ("diagnostic", "blueprint", "1 : N"),
                ("definition", "student_mis", "1 : N"),
                ("evidence", "diagnosis", "1 : N"),
                ("diagnosis", "candidate", "1 : N"),
                ("diagnosis", "verification", "1 : N"),
                ("verification", "session", "uses"),
                ("diagnosis", "repair", "1 : N"),
                ("repair", "repair_item", "1 : N"),
            ],
        ),
    ),
    (
        "05_management_exam_admission",
        graph(
            "Teaching, Assignment, Exam, Study Planning và Admission",
            [
                node("classroom", "classrooms", ["PK id", "name", "academic_year"]),
                node("plan", "teaching_plans", ["PK id", "FK classroom_id", "FK subject_id", "FK created_by", "status"], CREAM),
                node("plan_item", "teaching_plan_items", ["PK id", "FK teaching_plan_id", "FK problem_type_id", "FK learning_outcome_id"], CREAM),
                node("assignment", "assignments", ["PK id", "FK classroom_id", "FK subject_id", "FK teaching_plan_id?", "title", "due_at"], SKY),
                node("assignment_q", "assignment_questions", ["PK/FK assignment_id", "PK/FK question_id"], SKY),
                node("student_assignment", "student_assignments", ["PK id", "FK assignment_id", "FK student_id", "status", "score"], SKY),
                node("exam", "exams", ["PK id", "FK subject_id", "FK created_by", "title", "status"], "F0EAF9"),
                node("section", "exam_sections", ["PK id", "FK exam_id", "name", "order_index"], "F0EAF9"),
                node("exam_q", "exam_questions", ["PK id", "FK exam_id", "FK section_id?", "FK question_id"], "F0EAF9"),
                node("exam_attempt", "exam_attempts", ["PK id", "FK exam_id", "FK student_id", "score", "status"], "F0EAF9"),
                node("goal", "academic_goals", ["PK id", "FK student_id", "name", "target_date", "status"], "FDECC8"),
                node("subject_goal", "subject_goals", ["PK id", "FK academic_goal_id", "FK subject_id", "target_score"], "FDECC8"),
                node("study_plan", "study_plans", ["PK id", "FK student_id", "FK academic_goal_id?", "start_date", "end_date"], "FDECC8"),
                node("task", "study_tasks", ["PK id", "FK study_plan_id", "FK subject/concept/PT", "status"], "FDECC8"),
                node("university", "universities", ["PK id", "code", "name"], "DDEFF7"),
                node("major", "majors", ["PK id", "FK university_id", "code", "name"], "DDEFF7"),
                node("target", "student_admission_targets", ["PK id", "FK student_id", "FK university_id", "FK major_id"], "DDEFF7"),
            ],
            [
                ("classroom", "plan", "1 : N"),
                ("plan", "plan_item", "1 : N"),
                ("classroom", "assignment", "1 : N"),
                ("plan", "assignment", "1 : N"),
                ("assignment", "assignment_q", "1 : N"),
                ("assignment", "student_assignment", "1 : N"),
                ("exam", "section", "1 : N"),
                ("exam", "exam_q", "1 : N"),
                ("exam", "exam_attempt", "1 : N"),
                ("goal", "subject_goal", "1 : N"),
                ("goal", "study_plan", "1 : N"),
                ("study_plan", "task", "1 : N"),
                ("university", "major", "1 : N"),
                ("university", "target", "1 : N"),
                ("major", "target", "1 : N"),
            ],
        ).replace("rankdir=LR", "rankdir=TB"),
    ),
]


FLOW_DOT = """
digraph FLOW {
  graph [rankdir=LR, bgcolor="transparent", pad="0.2", nodesep="0.38", ranksep="0.5"];
  node [shape=box, style="rounded,filled", fontname="Arial", fontsize="10", color="#1E5B43", fillcolor="#DFF3E8", margin="0.14,0.09"];
  edge [fontname="Arial", fontsize="9", color="#1E5B43", fontcolor="#1E5B43", arrowsize="0.8"];
  curriculum [label="Curriculum +\nKnowledge Graph"];
  question [label="APPROVED\nQuestion Bank", fillcolor="#F0EAF9"];
  session [label="Practice / Diagnostic\nExam / Assignment", fillcolor="#FFF9ED"];
  attempt [label="Question Attempt"];
  evaluation [label="Deterministic\nEvaluation"];
  evidence [label="Learning Evidence", fillcolor="#DDEFF7"];
  mastery [label="Concept Mastery +\nProblemType Mastery", fillcolor="#DDF4D9"];
  twin [label="Student Digital Twin\nRead Model", fillcolor="#F6EAF2"];
  recommendation [label="Recommendation +\nReview + Repair", fillcolor="#FDECC8"];
  tree [label="Living Knowledge Tree", fillcolor="#DFF3E8"];
  curriculum -> question [label="classification"];
  question -> session [label="selection"];
  session -> attempt;
  attempt -> evaluation;
  evaluation -> evidence;
  evidence -> mastery;
  mastery -> twin;
  twin -> recommendation;
  curriculum -> tree [label="structure"];
  mastery -> tree [label="student state"];
}
"""


def render_dot(dot_text: str, output: Path) -> None:
    dot_path = output.with_suffix(".dot")
    dot_path.write_text(dot_text, encoding="utf-8")
    subprocess.run(
        ["dot", "-Tpng", "-Gdpi=180", str(dot_path), "-o", str(output)],
        check=True,
        capture_output=True,
        text=True,
    )


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.extend([field_begin, instruction, field_end])


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_fitted_picture(paragraph, image_path: Path) -> None:
    data = image_path.read_bytes()
    width_px, height_px = struct.unpack(">II", data[16:24])
    max_width = 10.2
    max_height = 5.85
    fitted_width = min(max_width, max_height * width_px / height_px)
    paragraph.add_run().add_picture(str(image_path), width=Inches(fitted_width))


def build_document(images: list[tuple[str, Path]], flow_image: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(297)
    section.page_height = Mm(210)
    section.top_margin = Mm(14)
    section.bottom_margin = Mm(14)
    section.left_margin = Mm(14)
    section.right_margin = Mm(14)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(30)
    styles["Title"].font.color.rgb = RGBColor(30, 91, 67)
    for style_name, size in (("Heading 1", 20), ("Heading 2", 15)):
        styles[style_name].font.name = "Aptos Display"
        styles[style_name].font.size = Pt(size)
        styles[style_name].font.color.rgb = RGBColor(30, 91, 67)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("EduMS — Sơ đồ ERD tổng quát")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run(
        "AI-Native Education Management & University Entrance Exam Support System"
    )
    subtitle_run.bold = True
    subtitle_run.font.size = Pt(15)
    subtitle_run.font.color.rgb = RGBColor(106, 90, 163)
    document.add_paragraph()
    intro = document.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.add_run(
        "Tài liệu được đối chiếu từ SQLAlchemy metadata hiện tại: 132 bảng vật lý. "
        "Các sơ đồ dưới đây tập trung vào bảng nghiệp vụ và quan hệ cốt lõi."
    )

    summary = document.add_table(rows=1, cols=4)
    summary.style = "Table Grid"
    headers = ["Phạm vi", "Cấu trúc cây", "Mastery", "Database"]
    values = [
        "Student · Teacher · Parent · Admin",
        "6 Chapter · 25 Topic · 75 ProblemType",
        "Concept + ProblemType",
        "PostgreSQL · 132 bảng",
    ]
    for index, text in enumerate(headers):
        summary.rows[0].cells[index].text = text
        set_cell_shading(summary.rows[0].cells[index], GREEN)
        for run in summary.rows[0].cells[index].paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
    row = summary.add_row()
    for index, text in enumerate(values):
        row.cells[index].text = text
        set_cell_shading(row.cells[index], "F6FBF8")

    document.add_paragraph()
    document.add_heading("Quy ước", level=1)
    add_bullets(
        document,
        [
            "PK: khóa chính; FK: khóa ngoại; UK: ràng buộc duy nhất.",
            "1 : N là quan hệ một-nhiều; 1 : 0..1 là quan hệ một-không hoặc một.",
            "Bảng mapping biểu diễn quan hệ nhiều-nhiều và có thể chứa weight/is_primary.",
            "Student Digital Twin là read model tổng hợp, không phải bảng giant JSON.",
            "Các dataset staging/review trên filesystem không được xem là dữ liệu production trong ERD.",
        ],
    )

    labels = [
        ("1. Identity, RBAC, Profile và Classroom", "Người dùng có đúng một role. Profile tách theo vai trò; quyền Teacher/Parent kiểm tra qua bảng quan hệ thật."),
        ("2. Curriculum, Knowledge Graph và Question Bank", "Cây học thuật đi từ CurriculumVersion đến ProblemType. Concept Graph chạy song song và được nối bằng các bảng mapping có trọng số."),
        ("3. Practice, Evidence, Mastery và Digital Twin", "Attempt tạo evidence sau khi được đánh giá. Mastery lưu riêng cho Concept và ProblemType; Chapter/Topic là aggregate khi đọc."),
        ("4. Diagnostic, Misconception và Adaptive Repair", "Diagnostic thu evidence; Root Cause phân tích ứng viên nguyên nhân; Repair tạo chuỗi bài sửa lỗi có kiểm chứng."),
        ("5. Teaching, Assignment, Exam, Planning và Admission", "Các lớp học và kế hoạch giảng dạy có thể giao bài, nhưng self-study vẫn hoạt động độc lập với Classroom."),
    ]
    for (file_name, image_path), (heading, description) in zip(images, labels, strict=True):
        document.add_page_break()
        document.add_heading(heading, level=1)
        document.add_paragraph(description)
        picture = document.add_paragraph()
        picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_fitted_picture(picture, image_path)

    document.add_page_break()
    document.add_heading("6. Luồng dữ liệu học tập tổng thể", level=1)
    document.add_paragraph(
        "Database là source of truth. AI chỉ hỗ trợ giải thích hoặc phân tích khi cần; "
        "chấm điểm cơ bản, evidence, mastery và recommendation nền tảng vẫn deterministic."
    )
    picture = document.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_fitted_picture(picture, flow_image)

    document.add_heading("Trạng thái dữ liệu Toán 12 tại thời điểm xuất", level=2)
    data_table = document.add_table(rows=1, cols=2)
    data_table.style = "Table Grid"
    for idx, title_text in enumerate(("Hạng mục", "Số lượng / trạng thái")):
        data_table.rows[0].cells[idx].text = title_text
        set_cell_shading(data_table.rows[0].cells[idx], GREEN)
        for run in data_table.rows[0].cells[idx].paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
    data_rows = [
        ("Chapter đang hoạt động", "6"),
        ("Topic đang hoạt động", "25"),
        ("ProblemType đang hoạt động", "75"),
        ("Concept trong read model", "89"),
        ("LearningOutcome", "34"),
        ("ConceptRelation", "93"),
        ("Question production trong database", "0 — cây xem được, Practice chưa có câu hỏi để phân phối"),
        ("Synthetic data", "Có marker riêng; production analytics có thể loại trừ"),
    ]
    for key, value in data_rows:
        cells = data_table.add_row().cells
        cells[0].text = key
        cells[1].text = value

    document.add_heading("Các domain phụ trợ không bung chi tiết trong sơ đồ", level=2)
    add_bullets(
        document,
        [
            "Question import, acquisition, validation và supply-gap audit.",
            "Teacher materials, material item taxonomy và bài làm tự luận theo từng bước.",
            "RAG documents/document_chunks và AI request logs.",
            "Tutor sessions, student learning memory và tutor summaries.",
            "Import history, provenance, review overlay và visual assets trên filesystem.",
        ],
    )

    for current_section in document.sections:
        add_page_number(current_section.footer.paragraphs[0])

    document.core_properties.title = "EduMS — Sơ đồ ERD tổng quát"
    document.core_properties.subject = "Kiến trúc dữ liệu tổng quát của Coach Learning / EduMS"
    document.core_properties.author = "EduMS Development Team"
    document.core_properties.keywords = "EduMS, Coach Learning, ERD, PostgreSQL, SQLAlchemy"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="edums_erd_") as temp_name:
        temp = Path(temp_name)
        images: list[tuple[str, Path]] = []
        for name, dot_text in DIAGRAMS:
            image_path = temp / f"{name}.png"
            render_dot(dot_text, image_path)
            images.append((name, image_path))
        flow_image = temp / "06_learning_flow.png"
        render_dot(FLOW_DOT, flow_image)
        build_document(images, flow_image)
    print(OUTPUT)


if __name__ == "__main__":
    main()
